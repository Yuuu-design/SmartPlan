"""Excel 导入接口测试：模板下载 + multipart 上传排产 + 脏数据 400。"""

from __future__ import annotations

import io

import pandas as pd
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from src.api.v1.schedule import app
from src.core.config import CONFIG
from src.data.loader import COL_DUE, COL_END, COL_ORDER_ID, COL_QTY_M, COL_SPEC

client = TestClient(app)


def _orders_xlsx(rows: list[dict]) -> bytes:
    df = pd.DataFrame(rows)
    buf = io.BytesIO()
    df.to_excel(buf, index=False)
    return buf.getvalue()


def _valid_rows(n: int = 2) -> list[dict]:
    # 从真实订单表取前 n 个未结束订单，保证规格能命中内置产能表
    df = pd.read_excel(CONFIG.order_file)
    df = df[~df[COL_END].isin(["已结束", "指定结束"])]
    df = df[df[COL_SPEC].notna() & (df[COL_QTY_M] > 0)].head(n)
    return df.to_dict("records")


def test_template_download_is_valid_xlsx_with_headers():
    res = client.get("/api/v1/schedule/template/orders")
    assert res.status_code == 200
    assert "spreadsheetml" in res.headers["content-type"]
    assert "attachment" in res.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(res.content))
    ws = wb["订单导入"]
    headers = [c.value for c in ws[1]]
    assert COL_ORDER_ID in headers
    assert COL_SPEC in headers
    assert COL_QTY_M in headers
    assert COL_DUE in headers
    assert "填写说明" in wb.sheetnames


def test_upload_orders_returns_schedule():
    content = _orders_xlsx(_valid_rows(2))
    res = client.post(
        "/api/v1/schedule/run",
        files={"order_file": ("orders.xlsx", content,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"limit_orders": "2"},
    )
    assert res.status_code == 200, res.text
    body = res.json()
    assert len(body["scheduled_tasks"]) == 6  # 2 单 × 3 工序
    assert body["kpis"] is not None


def test_upload_bad_excel_returns_400_with_reason():
    # 缺少必需列“规格”
    content = _orders_xlsx([{COL_ORDER_ID: "X1", COL_QTY_M: 100, COL_DUE: "2026-09-25", COL_END: ""}])
    res = client.post(
        "/api/v1/schedule/run",
        files={"order_file": ("bad.xlsx", content,
                              "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 400
    assert "缺少必需列" in res.json()["detail"]


def test_upload_missing_file_field_returns_400():
    # multipart 请求但未带 order_file 字段
    res = client.post(
        "/api/v1/schedule/run",
        files={"other": ("a.xlsx", b"x",
                         "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert res.status_code == 400
    assert "order_file" in res.json()["detail"]
